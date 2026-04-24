import os
import io
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image


def create_test_report_docx(data: dict) -> bytes:
    doc = Document()
    
    title = doc.add_heading('RAPPORT DE TEST - AI QA AGENT', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Info header
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    target_url = data.get('target_url', data.get('url', 'N/A'))
    browser = data.get('browser', 'N/A')
    page_info = data.get('page_info', {})
    lang = page_info.get('language', 'N/A')
    form_type = page_info.get('form_type', 'N/A')
    created_creds = data.get('created_credentials', {})
    creds_email = created_creds.get('email', 'N/A')
    
    # Metadata table
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Table Grid'
    
    meta_data = [
        ('Date', now),
        ('URL Testée', target_url),
        ('Browser', browser),
        ('Langue', lang),
        ('Type de Formulaire', form_type),
        ('Compte Créé', creds_email),
        ('Status', 'Terminé')
    ]
    
    for i, (key, val) in enumerate(meta_data):
        table.rows[i].cells[0].text = key
        table.rows[i].cells[1].text = str(val)
    
    doc.add_paragraph()
    
    # Summary section
    doc.add_heading('RÉSUMÉ', level=1)
    total = data.get('total_tests', len(data.get('results', [])))
    passed = data.get('passed', 0)
    failed = data.get('failed', 0)
    
    summary_text = f"Total: {total} tests\nPassed: {passed}\nFailed: {failed}"
    p = doc.add_paragraph(summary_text)
    p.runs[0].font.size = Pt(14)
    
    # Calculate pass rate
    pass_rate = (passed / total * 100) if total > 0 else 0
    p = doc.add_paragraph(f"Taux de réussite: {pass_rate:.1f}%")
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(34, 197, 94) if pass_rate >= 70 else RGBColor(239, 68, 68)
    
    doc.add_paragraph()
    
    # Test details
    doc.add_heading('DÉTAIL DES TESTS', level=1)
    
    results = data.get('results', [])
    for i, res in enumerate(results):
        tc = res.get('test_case', {})
        tc_id = tc.get('id', f'Test {i+1}')
        tc_desc = tc.get('description', '')
        status = res.get('status', 'unknown')
        
        # Test header
        doc.add_heading(f'{tc_id}: {tc_desc}', level=2)
        
        # Steps
        steps = tc.get('steps', [])
        if steps:
            doc.add_paragraph('Étapes:', style='Intense Quote')
            for step_idx, step in enumerate(steps, 1):
                field = step.get('field', '')
                value = step.get('value', '')
                # Mask password
                if 'password' in field.lower():
                    value = '******'
                p = doc.add_paragraph(f'  {step_idx}. {field} = {value}')
        
        # Result
        result_para = doc.add_paragraph()
        result_para.add_run('Résultat: ').bold = True
        if status == 'passed':
            result_para.add_run('✅ PASSED').font.color.rgb = RGBColor(34, 197, 94)
        else:
            result_para.add_run('❌ FAILED').font.color.rgb = RGBColor(239, 68, 68)
        
        # Error message
        if res.get('error'):
            p = doc.add_paragraph(f"Erreur: {res.get('error')}")
            p.font.color.rgb = RGBColor(239, 68, 68)
        
        # Screenshots
        screenshots = res.get('screenshots', [])
        if screenshots:
            doc.add_paragraph('Screenshots:')
            for shot in screenshots:
                shot_path = shot.get('path', '')
                if shot_path and os.path.exists(shot_path):
                    try:
                        # Resize image to max 400px width
                        img = Image.open(shot_path)
                        img.thumbnail((400, 300))
                        resized_path = shot_path.replace('.png', '_resized.png')
                        img.save(resized_path)
                        
                        doc.add_picture(resized_path, width=Inches(3.5))
                        last_para = doc.paragraphs[-1]
                        last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        
                        # Clean up resized image
                        try:
                            os.remove(resized_path)
                        except:
                            pass
                    except Exception as e:
                        p = doc.add_paragraph(f'[Image non disponible: {shot_path}]')
        
        doc.add_paragraph()
    
    # Footer
    doc.add_paragraph()
    p = doc.add_paragraph('— Document généré par AI QA Agent —')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    
    # Save to bytes
    file_buffer = io.BytesIO()
    doc.save(file_buffer)
    file_buffer.seek(0)
    
    return file_buffer.getvalue()


def save_report_docx(data: dict, output_path: str):
    doc_bytes = create_test_report_docx(data)
    with open(output_path, 'wb') as f:
        f.write(doc_bytes)
    return output_path